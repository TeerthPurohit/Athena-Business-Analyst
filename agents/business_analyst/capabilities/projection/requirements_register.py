"""Evidence-linked Excel register for every current detailed requirement."""

from collections import defaultdict
from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import PurePosixPath
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from agents.business_analyst.capabilities.projection import plain, requirements
from agents.business_analyst.capabilities.projection.business_analysis import (
    get_business_analysis,
    numbered_business_requirements,
)
from agents.business_analyst.capabilities.projection.document_template import project_document_context
from agents.business_analyst.facts import BATenantContext, get_facts
from agents.business_analyst.models import BaDeliverableSpec, BaSource
from agents.business_analyst.capabilities.projection import RenderedArtifact

TITLE = "Requirements Register"
NEEDS = "current requirement records"
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_REQUIRED_FIELDS = (
    ("Requirement type", "requirement_type"),
    ("Stakeholder", "stakeholder"),
    ("Owner", "owner"),
    ("Business problem", "business_problem"),
    ("Business objective", "business_objective"),
    ("Business benefit", "benefit"),
    ("Priority", "priority"),
    ("Acceptance criteria", "acceptance_criteria"),
)
_TECHNICAL_TERMS = re.compile(
    r"\b(JSONL|SQLite|tree-sitter|LSP|MCP|OAuth|PKCE|Python|Rust|Docker|keyring|"
    r"provider adapter|CLI command|database schema|API endpoint|container image)\b",
    re.IGNORECASE,
)


def _string(value):
    if value is None:
        return ""
    if isinstance(value, (list, tuple, set)):
        return "; ".join(text for item in value if (text := _string(item)))
    if isinstance(value, dict):
        return "; ".join(f"{key}: {text}" for key, item in value.items() if (text := _string(item)))
    return plain(value)


def _category(value):
    return _string(value.get("requirement_type") or value.get("category")).strip().lower().replace("-", "_") or "unclassified"


def _items(value):
    if value is None:
        return []
    values = value if isinstance(value, (list, tuple, set)) else [value]
    return list(dict.fromkeys(text for item in values if (text := _string(item))))


def _known_acceptance(value):
    explicit = _items(value.get("acceptance_criteria"))
    if explicit:
        return "; ".join(explicit)
    parts = []
    if _string(value.get("preconditions")):
        parts.append("Given " + _string(value["preconditions"]))
    if _string(value.get("trigger")):
        parts.append("When " + _string(value["trigger"]))
    if _string(value.get("outcomes")):
        parts.append("Then " + _string(value["outcomes"]))
    return "; ".join(parts)


def _quality(value, has_trace, has_business_link):
    missing = []
    for display, field in _REQUIRED_FIELDS:
        candidate = value.get(field)
        if field == "requirement_type":
            candidate = candidate or value.get("category")
        if field == "acceptance_criteria":
            candidate = candidate or (_items(value.get("trigger")) and _items(value.get("outcomes")))
        if not _string(candidate):
            missing.append(display)
    if not _string(value.get("task")) and not _string(value.get("object")):
        missing.append("Clear requirement statement")
    for field in ("dependencies", "assumptions", "constraints"):
        if not _string(value.get(field)):
            missing.append(field.replace("_", " ").capitalize())
    if not has_trace:
        missing.append("Source traceability")
    if not has_business_link:
        missing.append("Traceability to a business requirement")
    if _category(value) in {"business", "business_requirement"} and not _string(value.get("success_metric")):
        missing.append("Success metric where applicable")
    missing = list(dict.fromkeys(missing))
    status = "Draft — needs clarification" if missing else "Draft — review required"
    return status, missing


def _safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (int, float, bool)):
        return value
    text = str(value)
    if len(text) > 32000:
        text = text[:31980] + " … [cell text truncated; see linked source]"
    if text.lstrip().startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text


def _row_height(row, widths):
    max_lines = 1
    for value, width in zip(row, widths):
        if not isinstance(value, str) or not value:
            continue
        capacity = max(12, int(width * 1.25))
        lines = sum(
            max(1, (len(line) + capacity - 1) // capacity)
            for line in value.splitlines() or [value]
        )
        max_lines = max(max_lines, lines)
    return min(120, max(24, max_lines * 13 + 8))


def _add_table_sheet(workbook, title, headers, rows, widths, table_name):
    sheet = workbook.create_sheet(title)
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.append([_safe_cell(value) for value in headers])
    for row in rows:
        sheet.append([_safe_cell(value) for value in row])
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="4C365C")
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for row_index in range(2, sheet.max_row + 1):
        row_values = [sheet.cell(row_index, col).value for col in range(1, sheet.max_column + 1)]
        sheet.row_dimensions[row_index].height = _row_height(row_values, widths)
        for cell in sheet[row_index]:
            cell.font = Font(name="Aptos", size=10, color="2D2730")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        status_col = next((i + 1 for i, header in enumerate(headers) if header == "Quality Status"), None)
        if status_col:
            status_cell = sheet.cell(row_index, status_col)
            color = "FCE4D6" if "needs clarification" in str(status_cell.value).lower() else "FFF2CC"
            status_cell.fill = PatternFill("solid", fgColor=color)
    if sheet.max_row >= 2:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = sheet.dimensions
    return sheet


def _source_label(source):
    ref = str(source.ref or "").replace("\\", "/")
    title = PurePosixPath(ref).name if ref else "Recorded project context"
    return title, source.kind or "unknown", ref


def _technical_reason(value):
    requirement_type = _category(value)
    statement = " ".join(_string(value.get(key)) for key in ("task", "object", "comments", "data_format"))
    if requirement_type == "architecture_decision":
        return "Explicit architecture decision in the source record"
    if _TECHNICAL_TERMS.search(statement):
        return "Potential technical-design detail; confirm classification"
    return ""


async def render(spec: BaDeliverableSpec, ctx: BATenantContext, session: AsyncSession) -> RenderedArtifact:
    facts = await get_facts(ctx, session)
    reqs = requirements(facts)
    name, summary, context, _ = await project_document_context(ctx, session)
    analysis = await get_business_analysis(ctx, session, name, summary, context, reqs)
    business_requirements = numbered_business_requirements(analysis)
    business_ids_by_source_req = defaultdict(list)
    for business_requirement in business_requirements:
        for source_id in business_requirement.get("source_requirement_ids", []):
            business_ids_by_source_req[source_id].append(business_requirement["id"])

    req_facts = {
        fact.subject_key: fact for fact in facts
        if fact.subject_type == "Requirement" and fact.predicate == "specified_as" and isinstance(fact.value, dict)
    }
    span_facts = {
        fact.subject_key: fact for fact in facts
        if fact.subject_type == "SourceSpan" and fact.predicate == "evidence" and isinstance(fact.value, dict)
    }
    traced_spans = defaultdict(list)
    for fact in facts:
        if fact.subject_type == "Requirement" and fact.predicate == "derived_from" and fact.object_type == "SourceSpan" and fact.object_key:
            traced_spans[fact.subject_key].append(fact.object_key)

    source_ids = {
        fact.source_id for fact in req_facts.values() if fact.source_id
    } | {
        fact.source_id for fact in span_facts.values() if fact.source_id
    }
    source_records = {}
    if source_ids:
        result = await session.execute(
            select(BaSource).where(
                BaSource.id.in_(source_ids),
                BaSource.org_id == ctx.org_id,
                BaSource.project_id == ctx.project_id,
            )
        )
        source_records = {source.id: source for source in result.scalars().all()}

    requirement_rows = []
    story_rows = []
    acceptance_rows = []
    traceability_rows = []
    technical_rows = []
    evidence_rows = {}
    source_id_to_requirement_ids = defaultdict(list)
    frd_ids_by_requirement = {}
    acceptance_ids_by_requirement = defaultdict(list)
    test_ids_by_requirement = defaultdict(list)
    quality_counts = defaultdict(int)
    frd_type_counts = defaultdict(int)
    for subject_key, requirement_id, value in reqs:
        fact = req_facts.get(subject_key)
        span_ids = list(dict.fromkeys(
            [str(span_id) for span_id in value.get("evidence_span_ids", []) if span_id]
            + traced_spans.get(subject_key, [])
        ))
        spans = [span_facts[span_id] for span_id in span_ids if span_id in span_facts]
        requirement_source_ids = list(dict.fromkeys(
            ([fact.source_id] if fact and fact.source_id else [])
            + [span.source_id for span in spans if span.source_id]
        ))
        source_labels = []
        for source_id in requirement_source_ids:
            source = source_records.get(source_id)
            if not source:
                continue
            title, kind, ref = _source_label(source)
            source_labels.append(f"{title} ({kind})" + (f" — {ref}" if ref else ""))
            source_id_to_requirement_ids[source_id].append(requirement_id)
        source_labels = list(dict.fromkeys(source_labels))

        requirement_type = _category(value)
        frd_prefix = {
            "functional": "FR", "nonfunctional": "NFR", "business_rule": "BRULE",
            "constraint": "CST", "assumption": "ASM", "transitional": "TR",
        }.get(requirement_type)
        if frd_prefix:
            frd_type_counts[requirement_type] += 1
            frd_id = f"{frd_prefix}-{frd_type_counts[requirement_type]:03d}"
        else:
            frd_id = ""
        frd_ids_by_requirement[requirement_id] = frd_id
        business_ids = business_ids_by_source_req.get(requirement_id, [])
        quality_status, missing = _quality(
            value, bool(span_ids or requirement_source_ids), bool(business_ids)
        )
        quality_counts[quality_status] += 1
        task = _string(value.get("task"))
        obj = _string(value.get("object"))
        statement = task
        if obj and obj.casefold() not in statement.casefold():
            statement = f"{statement} — {obj}" if statement else obj
        if not statement:
            statement = _string(value.get("outcomes"))
        stakeholder = _string(value.get("stakeholder"))
        benefit = _string(value.get("benefit"))
        story_missing = [
            heading for heading, present in (
                ("Stakeholder", stakeholder), ("Task", task or obj), ("Business benefit", benefit),
            ) if not present
        ]
        if not story_missing:
            role = stakeholder.strip()
            if not role.lower().startswith(("a ", "an ", "the ")):
                article = "an" if role[:1].lower() in "aeiou" else "a"
                role = f"{article} {role}"
            story_text = f"As {role}, I want {task or obj}"
            if obj and obj.casefold() not in (task or "").casefold():
                story_text += f" ({obj})"
            story_text += f", so that {benefit}."
        else:
            story_text = ""
        story_rows.append((
            f"US-{requirement_id}", ", ".join(business_ids), story_text, ", ".join(story_missing),
            ", ".join(span_ids), "Draft — needs clarification" if story_missing else "Draft — review required",
        ))

        explicit_criteria = _items(value.get("acceptance_criteria"))
        if explicit_criteria:
            for criterion_index, criterion in enumerate(explicit_criteria, 1):
                criterion_id = f"AC-{requirement_id}-{criterion_index:02d}"
                acceptance_ids_by_requirement[requirement_id].append(criterion_id)
                test_ids_by_requirement[requirement_id].append(f"TC-{requirement_id}-{criterion_index:02d}")
                acceptance_rows.append((
                    criterion_id, requirement_id,
                    ", ".join(business_ids), criterion, "Source-authored",
                    ", ".join(span_ids), "Draft — review required",
                ))
        else:
            given = _items(value.get("preconditions"))
            when = _items(value.get("trigger"))
            then = _items(value.get("outcomes"))
            if when and then:
                scenario = " | ".join(
                    (["Given " + "; ".join(given)] if given else [])
                    + ["When " + "; ".join(when), "Then " + "; ".join(then)]
                )
                origin = "Derived only from recorded preconditions, trigger, and outcomes"
                criterion_status = "Draft — review required"
            else:
                scenario = ""
                origin = "No evidenced acceptance behavior"
                criterion_status = "Draft — needs clarification"
            criterion_id = f"AC-{requirement_id}-01"
            acceptance_ids_by_requirement[requirement_id].append(criterion_id)
            if when and then:
                test_ids_by_requirement[requirement_id].append(f"TC-{requirement_id}")
            acceptance_rows.append((
                criterion_id, requirement_id, ", ".join(business_ids),
                scenario, origin, ", ".join(span_ids), criterion_status,
            ))

        excerpts = []
        for span_id in span_ids:
            span_fact = span_facts.get(span_id)
            if not span_fact:
                continue
            span_value = span_fact.value or {}
            excerpt = _string(span_value.get("text"))
            excerpts.append(f"{span_id}: {excerpt}")
            source = source_records.get(span_fact.source_id)
            evidence_rows[(span_id, requirement_id)] = (
                span_id, requirement_id, _source_label(source)[0] if source else "",
                source.kind if source else "", source.ref if source else "",
                span_value.get("ordinal"), span_value.get("start_char"), span_value.get("end_char"),
                span_value.get("speaker"), span_value.get("timestamp"), excerpt,
            )

        requirement_rows.append((
            requirement_id, frd_id, f"US-{requirement_id}", requirement_type, _string(value.get("category")),
            _string(value.get("functional_area")), statement, _string(value.get("stakeholder")),
            _string(value.get("owner")), _string(value.get("business_problem")),
            _string(value.get("business_objective")), _string(value.get("benefit")),
            _string(value.get("success_metric")), _string(value.get("priority")),
            _string(value.get("status")), _known_acceptance(value), _string(value.get("trigger")),
            _string(value.get("preconditions")), _string(value.get("outcomes")),
            _string(value.get("dependencies")), _string(value.get("assumptions")),
            _string(value.get("business_rules")), _string(value.get("constraints")),
            _string(value.get("input_data")), _string(value.get("output_data")),
            _string(value.get("data_format")), _string(value.get("exceptions")),
            _string(value.get("comments")), ", ".join(business_ids), "\n".join(source_labels),
            ", ".join(span_ids), "\n".join(excerpts), "; ".join(missing), quality_status,
            fact.id if fact else "", fact.source_id if fact else "",
        ))
        objective_text = ", ".join(dict.fromkeys(
            _string(next((item.get("business_objective") for item in business_requirements if item["id"] == business_id), None))
            for business_id in business_ids
            if _string(next((item.get("business_objective") for item in business_requirements if item["id"] == business_id), None))
        ))
        trace_gaps = list(missing)
        trace_gaps.append("Implementation / verification evidence not linked")
        traceability_rows.append((
            objective_text, ", ".join(business_ids), frd_id, f"US-{requirement_id}", requirement_id,
            ", ".join(acceptance_ids_by_requirement[requirement_id]),
            ", ".join(test_ids_by_requirement[requirement_id]), ", ".join(span_ids),
            "", "Draft — incomplete", "; ".join(dict.fromkeys(trace_gaps)),
        ))

        technical_reason = _technical_reason(value)
        if technical_reason:
            technical_rows.append((
                f"TD-{len(technical_rows) + 1:03d}", technical_reason, requirement_id,
                requirement_type, statement, _string(value.get("business_problem")),
                _string(value.get("business_objective")), _string(value.get("benefit")),
                _string(value.get("dependencies")), _string(value.get("constraints")),
                _string(value.get("outcomes")), _string(value.get("owner")),
                "\n".join(source_labels), "Draft — needs architecture / engineering review",
            ))

    business_rows = []
    for item in business_requirements:
        missing = [
            heading for heading, key in (
                ("Stakeholder", "stakeholder"), ("Business problem", "business_problem"),
                ("Business objective", "business_objective"), ("Business benefit", "business_benefit"),
                ("Success metric", "success_metric"), ("Priority", "priority"),
            ) if not _string(item.get(key))
        ]
        source_ids = item.get("source_requirement_ids", [])
        business_rows.append((
            item["id"], item.get("title"), item.get("statement"), item.get("stakeholder"),
            item.get("business_problem"), item.get("business_objective"), item.get("business_benefit"),
            item.get("success_metric"), item.get("priority"), ", ".join(source_ids),
            ", ".join(item.get("functional_areas", [])), "; ".join(item.get("clarification_questions", [])),
            "; ".join(missing), "Draft — needs clarification" if missing else "Draft — review required",
        ))

    source_rows = []
    for source_id, source in source_records.items():
        if not source_id_to_requirement_ids.get(source_id):
            continue
        title, kind, ref = _source_label(source)
        source_rows.append((
            source_id, title, kind, source.tier, ref,
            ", ".join(sorted(set(source_id_to_requirement_ids[source_id]))),
        ))
    source_rows.sort(key=lambda row: (row[1].casefold(), row[0]))
    evidence_rows = sorted(evidence_rows.values(), key=lambda row: (row[2].casefold(), row[5] or 0, row[0], row[1]))

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Overview"
    overview.sheet_view.showGridLines = False
    overview.merge_cells("A1:D1")
    overview["A1"] = "ATHENA · REQUIREMENTS REGISTER"
    overview["A1"].font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    overview["A1"].fill = PatternFill("solid", fgColor="4C365C")
    overview["A1"].alignment = Alignment(vertical="center")
    overview.row_dimensions[1].height = 34
    overview_data = [
        ("Project", name),
        ("Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Current detailed requirements", len(requirement_rows)),
        ("Business requirement groups", len(business_rows)),
        ("Story records", len(story_rows)),
        ("Acceptance criteria records", len(acceptance_rows)),
        ("Traceability records", len(traceability_rows)),
        ("Draft — needs clarification", quality_counts["Draft — needs clarification"]),
        ("Draft — review required", quality_counts["Draft — review required"]),
        ("Potential technical decisions / details", len(technical_rows)),
        ("Source records", len(source_rows)),
        ("Evidence spans", len(evidence_rows)),
    ]
    for row_index, (heading, value) in enumerate(overview_data, 3):
        overview.cell(row_index, 1, _safe_cell(heading))
        overview.cell(row_index, 2, _safe_cell(value))
        overview.cell(row_index, 1).font = Font(name="Aptos", bold=True, color="4C365C")
        overview.cell(row_index, 2).font = Font(name="Aptos", color="2D2730")
        overview.cell(row_index, 1).alignment = Alignment(vertical="top", wrap_text=True)
        overview.cell(row_index, 2).alignment = Alignment(vertical="top", wrap_text=True)
    overview["A16"] = "How to use this workbook"
    overview["A16"].font = Font(name="Aptos", bold=True, size=12, color="4C365C")
    overview.merge_cells("A17:D17")
    overview["A17"] = "Filter the Requirements Register by type, area, priority, business requirement, or quality status. Source IDs and evidence spans preserve the link to captured project material."
    overview["A17"].alignment = Alignment(vertical="top", wrap_text=True)
    overview.merge_cells("A18:D18")
    overview["A18"] = "All generated entries remain drafts. Missing decisions are listed in Missing Information; no stakeholder, benefit, KPI target, priority, or approval status is inferred."
    overview["A18"].alignment = Alignment(vertical="top", wrap_text=True)
    overview.column_dimensions["A"].width = 42
    overview.column_dimensions["B"].width = 90
    overview.column_dimensions["C"].width = 20
    overview.column_dimensions["D"].width = 20
    overview.row_dimensions[17].height = 40
    overview.row_dimensions[18].height = 40

    requirement_headers = [
        "Requirement ID", "FRD Requirement ID", "User Story ID", "Requirement Type", "Source Category",
        "Functional Area", "Requirement Statement",
        "Stakeholder", "Owner", "Business Problem", "Business Objective", "Business Benefit",
        "Success Metric", "Priority", "Source-reported Status", "Acceptance Criteria", "Trigger",
        "Preconditions", "Expected Outcomes", "Dependencies", "Assumptions", "Business Rules",
        "Constraints", "Inputs", "Outputs", "Data Format", "Exceptions", "Comments",
        "Business Requirement IDs", "Source Documents", "Evidence Span IDs", "Source Evidence Excerpts",
        "Missing Information", "Quality Status", "Source Fact ID", "Source ID",
    ]
    requirement_widths = [13, 20, 20, 19, 17, 26, 54, 24, 24, 40, 40, 36, 32, 14, 22, 48, 32, 38, 40, 34, 34, 38, 38, 30, 30, 24, 36, 36, 22, 50, 42, 70, 44, 30, 38, 38]
    _add_table_sheet(workbook, "Requirements Register", requirement_headers, requirement_rows, requirement_widths, "RequirementsTable")

    business_headers = [
        "Business Requirement ID", "Title", "Business Requirement", "Stakeholder", "Business Problem",
        "Business Objective", "Business Benefit", "Success Metric", "Priority", "Detailed Requirement IDs",
        "Functional Areas", "Open Decisions", "Missing Information", "Quality Status",
    ]
    _add_table_sheet(
        workbook, "Business Requirements", business_headers, business_rows,
        [20, 34, 70, 24, 42, 42, 38, 34, 14, 38, 30, 46, 42, 30], "BusinessRequirementsTable",
    )
    _add_table_sheet(
        workbook, "User Stories",
        ["Story ID", "Business Requirement IDs", "User Story", "Missing Information", "Evidence Span IDs", "Quality Status"],
        story_rows, [20, 28, 90, 42, 44, 30], "UserStoriesTable",
    )
    _add_table_sheet(
        workbook, "Acceptance Criteria",
        ["Acceptance Criteria ID", "Requirement ID", "Business Requirement IDs", "Given / When / Then", "Basis", "Evidence Span IDs", "Quality Status"],
        acceptance_rows, [26, 18, 28, 90, 58, 44, 30], "AcceptanceCriteriaTable",
    )
    _add_table_sheet(
        workbook, "Traceability",
        ["Business Objective", "Business Requirement ID", "FRD Requirement ID", "User Story ID",
         "Requirement ID", "Acceptance Criteria IDs", "Test Case IDs", "Source Evidence Span IDs",
         "Implementation / Verification Evidence", "Traceability Status", "Missing Link / Decision"],
        traceability_rows, [48, 24, 22, 22, 18, 42, 42, 48, 56, 28, 56], "TraceabilityTable",
    )
    technical_headers = [
        "Technical Detail ID", "Classification", "Source Requirement ID", "Source Type",
        "Evidence-backed Decision / Detail", "Business Problem", "Business Objective", "Business Benefit",
        "Dependencies", "Constraints", "Expected Consequences", "Owner", "Source Documents", "Review Status",
    ]
    _add_table_sheet(
        workbook, "Technical Decisions", technical_headers, technical_rows,
        [18, 44, 22, 24, 62, 40, 40, 36, 34, 34, 40, 24, 52, 46], "TechnicalDecisionsTable",
    )
    _add_table_sheet(
        workbook, "Sources",
        ["Source ID", "Source Title", "Source Type", "Evidence Tier", "Source Reference", "Requirement IDs"],
        source_rows, [38, 42, 20, 20, 68, 56], "SourcesTable",
    )
    _add_table_sheet(
        workbook, "Evidence",
        ["Evidence Span ID", "Requirement ID", "Source Title", "Source Type", "Source Reference",
         "Ordinal", "Start Character", "End Character", "Speaker", "Timestamp", "Evidence Text"],
        evidence_rows, [42, 18, 42, 20, 68, 12, 16, 16, 24, 24, 100], "EvidenceTable",
    )

    output = BytesIO()
    workbook.save(output)
    preview = (
        f"Excel workbook preview\n\nProject: {name}\n"
        f"Current detailed requirements: {len(requirement_rows)}\n"
        f"Business requirement groups: {len(business_rows)}\n"
        f"Needs clarification: {quality_counts['Draft — needs clarification']}\n"
        f"Technical decision candidates: {len(technical_rows)}\n\n"
        "Sheets: Overview, Requirements Register, Business Requirements, User Stories, Acceptance Criteria, Traceability, Technical Decisions, Sources, Evidence.\n"
        "Download the workbook to filter records and review source traceability."
    )
    return RenderedArtifact(
        content=output.getvalue(), output_format="xlsx", preview=preview,
        media_type=MEDIA_TYPE, extension="xlsx",
    )


def xlsx_preview(content: bytes) -> str:
    """Create a human-readable preview for the deliverable Read panel."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        overview = workbook["Overview"]
        project = overview["B3"].value or "Project"
        values = {overview.cell(row, 1).value: overview.cell(row, 2).value for row in range(3, 15)}
        sheets = ", ".join(workbook.sheetnames)
        workbook.close()
        return (
            f"Excel workbook preview\n\nProject: {project}\n"
            f"Current detailed requirements: {values.get('Current detailed requirements', 0)}\n"
            f"Business requirement groups: {values.get('Business requirement groups', 0)}\n"
            f"Needs clarification: {values.get('Draft — needs clarification', 0)}\n"
            f"Technical decision candidates: {values.get('Potential technical decisions / details', 0)}\n\n"
            f"Sheets: {sheets}. Download the workbook to filter records and review source traceability."
        )
    except Exception:
        return "This is an Excel requirements register. Download the workbook to filter all requirements and review their source evidence."
